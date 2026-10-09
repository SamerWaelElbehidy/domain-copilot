# ruff: noqa: E501
"""Builds teaching/slides/session.pptx from slides.md.

The Markdown file is the single source. This script only lays it out, so the
content diffs in pull requests and the deck can be regenerated at any time.

    pip install -r teaching/requirements.txt
    python teaching/slides/build_pptx.py

Markdown conventions (one slide per `---`):
    # Title            slide title (the first slide also uses `## subtitle`)
    - bullet           bullets; two-space indent for a sub-bullet
    | a | b |          a table (first row is the header, second row is the separator)
    ```                a code block
    > text             a highlighted statement
    ::: flow A | B     a left-to-right process diagram
    <!-- notes: ... -->  speaker notes (may span lines)
"""

from __future__ import annotations

import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
SRC = HERE / "slides.md"
OUT = HERE / "session.pptx"

INK = RGBColor(0x1C, 0x23, 0x30)
MUTED = RGBColor(0x5D, 0x66, 0x75)
ACCENT = RGBColor(0x0B, 0x5F, 0xFF)
PANEL = RGBColor(0xF1, 0xF3, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WIDTH, HEIGHT = 13.333, 7.5
LEFT, BODY_WIDTH = 0.7, 11.9
BOTTOM = 7.1


@dataclass
class Block:
    kind: str  # bullets | table | code | quote | flow
    lines: list[str] = field(default_factory=list)


@dataclass
class Slide:
    title: str = ""
    subtitle: str = ""
    blocks: list[Block] = field(default_factory=list)
    notes: str = ""


def parse(text: str) -> list[Slide]:
    slides: list[Slide] = []
    for raw in re.split(r"\n---\n", text.strip()):
        slide = Slide()
        notes = re.search(r"<!--\s*notes:(.*?)-->", raw, re.S)
        if notes:
            slide.notes = " ".join(notes.group(1).split())
            raw = raw.replace(notes.group(0), "")
        lines = raw.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i]
            if not line.strip():
                i += 1
            elif line.startswith("# "):
                slide.title = line[2:].strip()
                i += 1
            elif line.startswith("## "):
                slide.subtitle = line[3:].strip()
                i += 1
            elif line.startswith("```"):
                block = Block("code")
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    block.lines.append(lines[i])
                    i += 1
                i += 1
                slide.blocks.append(block)
            elif line.startswith("|"):
                block = Block("table")
                while i < len(lines) and lines[i].startswith("|"):
                    block.lines.append(lines[i])
                    i += 1
                slide.blocks.append(block)
            elif line.startswith("::: flow"):
                slide.blocks.append(Block("flow", [line[len("::: flow") :].strip()]))
                i += 1
            elif line.startswith("> "):
                slide.blocks.append(Block("quote", [line[2:].strip()]))
                i += 1
            elif line.lstrip().startswith("- "):
                block = Block("bullets")
                while i < len(lines) and lines[i].lstrip().startswith("- "):
                    block.lines.append(lines[i])
                    i += 1
                slide.blocks.append(block)
            else:  # a plain paragraph
                slide.blocks.append(Block("bullets", [f"- {line.strip()}"]))
                i += 1
        slides.append(slide)
    return slides


def runs(paragraph, text: str, size: int, color=INK, mono: bool = False) -> None:
    """Adds text with **bold** and `code` spans."""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**"):
            run.text, run.font.bold = part[2:-2], True
        elif part.startswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
        else:
            run.text = part
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if mono:
            run.font.name = "Consolas"


def wrapped_lines(text: str, size: int, width_in: float) -> int:
    chars_per_line = max(20, int(width_in * 72 / (size * 0.47)))
    return max(1, math.ceil(len(re.sub(r"[*`]", "", text)) / chars_per_line))


def textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.05)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    return box, frame


def block_height(block: Block, size: int) -> float:
    if block.kind == "bullets":
        lines = sum(wrapped_lines(b.strip()[2:], size, BODY_WIDTH - 0.5) for b in block.lines)
        return lines * size * 1.3 / 72 + len(block.lines) * 0.1 + 0.05
    if block.kind == "code":
        return len(block.lines) * 0.3 + 0.35
    if block.kind == "table":
        return (len(block.lines) - 1) * 0.45 + 0.2
    if block.kind == "quote":
        lines = wrapped_lines(block.lines[0], size + 2, BODY_WIDTH - 0.6)
        return lines * (size + 2) * 1.3 / 72 + 0.25
    return 1.3  # flow


def draw(slide, block: Block, y: float, size: int) -> float:
    if block.kind == "bullets":
        _, frame = textbox(slide, LEFT, y, BODY_WIDTH, block_height(block, size))
        first = True
        for raw in block.lines:
            level = (len(raw) - len(raw.lstrip())) // 2
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            marker = "• " if level == 0 else "– "
            paragraph.level = 0
            runs(paragraph, marker + raw.strip()[2:], size - 2 * level)
            paragraph.space_after = Pt(6)
            paragraph.left_indent = (
                Inches(0.35 * level) if hasattr(paragraph, "left_indent") else None
            )
    elif block.kind == "code":
        height = block_height(block, size)
        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(LEFT), Inches(y), Inches(BODY_WIDTH), Inches(height)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = PANEL
        panel.line.fill.background()
        frame = panel.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = Inches(0.25)
        for index, line in enumerate(block.lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.LEFT
            runs(paragraph, line or " ", 15, INK, mono=True)
    elif block.kind == "quote":
        height = block_height(block, size)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(LEFT), Inches(y), Inches(0.08), Inches(height - 0.1)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        _, frame = textbox(slide, LEFT + 0.25, y, BODY_WIDTH - 0.3, height)
        runs(frame.paragraphs[0], block.lines[0], size + 2, ACCENT)
        for run in frame.paragraphs[0].runs:
            run.font.italic = True
    elif block.kind == "table":
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in block.lines]
        rows = [r for i, r in enumerate(rows) if not (i == 1 and set("".join(r)) <= set("-: "))]
        height = block_height(block, size)
        table = slide.shapes.add_table(
            len(rows), len(rows[0]), Inches(LEFT), Inches(y), Inches(BODY_WIDTH), Inches(height)
        ).table
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = ""
                runs(cell.text_frame.paragraphs[0], value, size - 3, WHITE if r == 0 else INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT if r == 0 else (PANEL if r % 2 == 0 else WHITE)
    else:  # flow
        steps = [s.strip() for s in block.lines[0].split("|")]
        gap = 0.12
        width = (BODY_WIDTH - gap * (len(steps) - 1)) / len(steps)
        for index, step in enumerate(steps):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.PENTAGON if index < len(steps) - 1 else MSO_SHAPE.RECTANGLE,
                Inches(LEFT + index * (width + gap)),
                Inches(y),
                Inches(width),
                Inches(1.05),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = ACCENT
            shape.line.fill.background()
            frame = shape.text_frame
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            runs(paragraph, step, 15 if len(steps) > 4 else 17, WHITE)
            for run in paragraph.runs:
                run.font.bold = True
    return y + block_height(block, size) + 0.15


def layout(deck: Presentation, slide_data: Slide, index: int, total: int) -> list[str]:
    warnings: list[str] = []
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if index == 0:
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(WIDTH), Inches(HEIGHT))
        band.fill.solid()
        band.fill.fore_color.rgb = INK
        band.line.fill.background()
        _, frame = textbox(slide, 0.9, 2.1, 11.5, 1.6)
        runs(frame.paragraphs[0], slide_data.title, 48, WHITE)
        for run in frame.paragraphs[0].runs:
            run.font.bold = True
        _, frame = textbox(slide, 0.9, 3.7, 11.5, 1.0)
        runs(frame.paragraphs[0], slide_data.subtitle, 28, RGBColor(0xC9, 0xD6, 0xFF))
        text = " ".join(b.lines[0].strip()[2:] for b in slide_data.blocks if b.lines)
        _, frame = textbox(slide, 0.9, 5.0, 11.5, 1.0)
        runs(frame.paragraphs[0], text, 18, RGBColor(0xC9, 0xCF, 0xDA))
    else:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), Inches(HEIGHT))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        _, frame = textbox(slide, LEFT, 0.35, BODY_WIDTH, 0.9)
        runs(frame.paragraphs[0], slide_data.title, 32, INK)
        for run in frame.paragraphs[0].runs:
            run.font.bold = True
        size = 22
        while (
            size > 15
            and sum(block_height(b, size) + 0.15 for b in slide_data.blocks) > BOTTOM - 1.4
        ):
            size -= 1
        y = 1.4
        for block in slide_data.blocks:
            y = draw(slide, block, y, size)
        if y > BOTTOM + 0.2:
            warnings.append(f"slide {index + 1} ('{slide_data.title}') may overflow (y={y:.1f})")
        _, frame = textbox(slide, WIDTH - 1.6, 7.05, 1.2, 0.35)
        runs(frame.paragraphs[0], f"{index + 1} / {total}", 11, MUTED)
        frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    if slide_data.notes:
        slide.notes_slide.notes_text_frame.text = slide_data.notes
    return warnings


def main() -> int:
    slides = parse(SRC.read_text(encoding="utf-8"))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    warnings: list[str] = []
    for index, slide_data in enumerate(slides):
        warnings += layout(deck, slide_data, index, len(slides))
    deck.save(OUT)
    print(
        f"wrote {OUT.name}: {len(slides)} slides, {sum(1 for s in slides if s.notes)} with speaker notes"
    )
    for warning in warnings:
        print("warning:", warning)
    return 0


if __name__ == "__main__":
    sys.exit(main())
